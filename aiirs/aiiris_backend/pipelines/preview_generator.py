import os
import base64
import io
import tempfile
from pathlib import Path
from typing import Dict, Any

try:
    import fitz  # PyMuPDF for PDF handling
    FITZ_AVAILABLE = True
except ImportError:
    FITZ_AVAILABLE = False

try:
    from PIL import Image, ImageDraw, ImageFont
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from docx2pdf import convert
    DOCX2PDF_AVAILABLE = True
except ImportError:
    DOCX2PDF_AVAILABLE = False


class PreviewGenerator:
    """
    Generates previews for various file types including PDF, images, Excel, DOCX, and text files.
    """
    
    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        self.file_extension = self.file_path.suffix.lower()
    
    def generate_preview(self) -> Dict[str, Any]:
        """
        Generate a preview based on the file type.
        Returns a dictionary with preview type and data.
        """
        if not self.file_path.exists():
            raise FileNotFoundError(f"File not found: {self.file_path}")
        
        try:
            if self.file_extension == '.pdf':
                return self._generate_pdf_preview()
            elif self.file_extension in ['.png', '.jpg', '.jpeg', '.gif', '.bmp']:
                return self._generate_image_preview()
            elif self.file_extension in ['.xlsx', '.xls']:
                return self._generate_excel_preview()
            elif self.file_extension in ['.docx', '.doc']:
                return self._generate_docx_preview()
            elif self.file_extension == '.txt':
                return self._generate_text_preview()
            else:
                return self._generate_default_preview()
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating preview: {str(e)}"
            }
    
    def _generate_pdf_preview(self) -> Dict[str, Any]:
        """Generate preview for PDF files - converts first page to image."""
        if not FITZ_AVAILABLE:
            return {
                "type": "error",
                "data": "PyMuPDF (fitz) library not available. Cannot generate PDF preview."
            }
        
        if not PIL_AVAILABLE:
            return {
                "type": "error",
                "data": "PIL (Pillow) library not available. Cannot generate PDF preview."
            }
        
        pdf_document = None
        try:
            pdf_document = fitz.open(str(self.file_path))
            
            if len(pdf_document) == 0:
                return {
                    "type": "error",
                    "data": "PDF file is empty"
                }
            
            # Store page count before closing
            page_count = len(pdf_document)
            
            # Get the first page
            first_page = pdf_document[0]
            
            # Convert page to image with good quality
            mat = fitz.Matrix(2.0, 2.0)  # Scale factor for better quality
            pix = first_page.get_pixmap(matrix=mat)
            
            # Convert to PIL Image
            img_data = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_data))
            
            # Resize for preview (max 400px width while maintaining aspect ratio)
            img = self._resize_image_for_preview(img, max_width=400)
            
            # Convert to base64
            buffer = io.BytesIO()
            img.save(buffer, format='PNG')
            img_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            return {
                "type": "image",
                "data": f"data:image/png;base64,{img_base64}",
                "metadata": {
                    "pages": page_count,
                    "title": "PDF Preview - First Page"
                }
            }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating PDF preview: {str(e)}"
            }
        finally:
            # Ensure document is closed even if an error occurs
            if pdf_document is not None:
                try:
                    pdf_document.close()
                except:
                    pass
    
    def _generate_image_preview(self) -> Dict[str, Any]:
        """Generate preview for image files."""
        if not PIL_AVAILABLE:
            return {
                "type": "error",
                "data": "PIL (Pillow) library not available. Cannot generate image preview."
            }
        
        try:
            img = Image.open(self.file_path)
            
            # Resize for preview
            img = self._resize_image_for_preview(img, max_width=400)
            
            # Convert to RGB if necessary (for JPEG compatibility)
            if img.mode in ('RGBA', 'LA', 'P'):
                img = img.convert('RGB')
            
            # Convert to base64
            buffer = io.BytesIO()
            format_type = 'JPEG' if self.file_extension in ['.jpg', '.jpeg'] else 'PNG'
            img.save(buffer, format=format_type, quality=85)
            img_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            mime_type = f"image/{'jpeg' if format_type == 'JPEG' else 'png'}"
            
            return {
                "type": "image",
                "data": f"data:{mime_type};base64,{img_base64}",
                "metadata": {
                    "format": format_type,
                    "size": f"{img.width}x{img.height}"
                }
            }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating image preview: {str(e)}"
            }
    
    def _generate_excel_preview(self) -> Dict[str, Any]:
        """Generate preview for Excel files - show first few rows of first sheet."""
        if not PANDAS_AVAILABLE:
            return {
                "type": "error",
                "data": "Pandas library not available. Cannot generate Excel preview."
            }
        
        try:
            # Read the first sheet with limited rows
            df = pd.read_excel(self.file_path, nrows=10)
            
            if df.empty:
                return {
                    "type": "text",
                    "data": "Excel file is empty or has no readable data."
                }
            
            # Convert to HTML table for preview
            html_table = df.to_html(classes='excel-preview', table_id='excel-preview-table', escape=False)
            
            # Also provide text summary
            summary = f"Excel file with {len(df.columns)} columns and {len(df)} rows (showing first 10 rows)\n\n"
            summary += f"Columns: {', '.join(df.columns.astype(str))}\n\n"
            summary += df.to_string(max_rows=10, max_cols=10)
            
            return {
                "type": "excel",
                "data": {
                    "html": html_table,
                    "text": summary,
                    "columns": list(df.columns.astype(str)),
                    "rows_shown": len(df),
                    "total_columns": len(df.columns)
                }
            }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating Excel preview: {str(e)}"
            }
    
    def _generate_docx_preview(self) -> Dict[str, Any]:
        """Generate preview for DOCX files - converts first page to image when possible."""
        if not DOCX_AVAILABLE:
            return {
                "type": "error",
                "data": "python-docx library not available. Cannot generate DOCX preview."
            }
        
        # Try to convert to PDF first, then to image
        if DOCX2PDF_AVAILABLE and FITZ_AVAILABLE and PIL_AVAILABLE:
            try:
                print(f"Attempting PDF conversion for: {self.file_path.name}")
                result = self._generate_docx_image_preview()
                print(f"PDF conversion successful for: {self.file_path.name}")
                return result
            except Exception as e:
                print(f"DOCX to image conversion failed for {self.file_path.name}: {e}, falling back to visual preview")
        
        # Fall back to creating a visual text representation
        if PIL_AVAILABLE:
            try:
                return self._generate_docx_visual_preview()
            except Exception as e:
                print(f"DOCX visual preview failed: {e}, falling back to simple text preview")
        
        # Final fallback: simple text preview
        return self._generate_docx_text_preview()
    
    def _generate_docx_image_preview(self) -> Dict[str, Any]:
        """Try to convert DOCX to PDF, then to image with retry logic."""
        return self._convert_docx_with_retry(max_retries=2)
    
    def _convert_docx_with_retry(self, max_retries: int = 2) -> Dict[str, Any]:
        """Convert DOCX with retry logic for COM errors."""
        for attempt in range(max_retries + 1):
            temp_pdf = None
            try:
                # Initialize COM for Windows (fresh initialization each attempt)
                import sys
                if sys.platform.startswith('win'):
                    try:
                        import pythoncom
                        # Try to uninitialize first in case there's a stale COM state
                        try:
                            pythoncom.CoUninitialize()
                        except:
                            pass
                        # Fresh initialization
                        pythoncom.CoInitialize()
                        print(f"COM initialized successfully for attempt {attempt + 1}")
                    except Exception as com_error:
                        print(f"COM initialization failed on attempt {attempt + 1}: {com_error}")
                        if attempt == max_retries:
                            raise Exception(f"COM initialization failed after {max_retries + 1} attempts")
                
                # Create a temporary PDF file
                with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as temp_file:
                    temp_pdf = temp_file.name
                
                print(f"Converting {self.file_path.name} to PDF (attempt {attempt + 1})")
                
                # Convert DOCX to PDF
                convert(str(self.file_path), temp_pdf)
                
                # Verify the PDF was created and is not empty
                if not os.path.exists(temp_pdf) or os.path.getsize(temp_pdf) == 0:
                    raise Exception("Generated PDF is empty or doesn't exist")
                
                print(f"PDF conversion successful for {self.file_path.name} on attempt {attempt + 1}")
                
                # Use our PDF preview logic
                return self._convert_pdf_to_image(temp_pdf)
                
            except Exception as e:
                print(f"Conversion attempt {attempt + 1} failed for {self.file_path.name}: {e}")
                
                # Clean up COM for this attempt
                import sys
                if sys.platform.startswith('win'):
                    try:
                        import pythoncom
                        pythoncom.CoUninitialize()
                    except:
                        pass
                
                # Clean up temporary PDF file
                if temp_pdf and os.path.exists(temp_pdf):
                    try:
                        os.unlink(temp_pdf)
                    except:
                        pass
                
                # If this was the last attempt, raise the exception
                if attempt == max_retries:
                    raise e
                
                # Wait a bit before retrying
                import time
                time.sleep(0.5)
        
        # This should never be reached, but just in case
        raise Exception("Unexpected error in conversion retry logic")
    
    def _convert_pdf_to_image(self, pdf_path: str) -> Dict[str, Any]:
        """Convert PDF to image preview."""
        pdf_document = None
        try:
            pdf_document = fitz.open(pdf_path)
            
            if len(pdf_document) == 0:
                raise Exception("Generated PDF is empty")
            
            # Get the first page
            first_page = pdf_document[0]
            
            # Convert page to image
            mat = fitz.Matrix(2.0, 2.0)
            pix = first_page.get_pixmap(matrix=mat)
            
            # Convert to PIL Image
            img_data = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_data))
            
            # Resize for preview
            img = self._resize_image_for_preview(img, max_width=400)
            
            # Convert to base64
            buffer = io.BytesIO()
            img.save(buffer, format='PNG')
            img_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            return {
                "type": "image",
                "data": f"data:image/png;base64,{img_base64}",
                "metadata": {
                    "title": "DOCX Preview - First Page",
                    "method": "pdf_conversion"
                }
            }
            
        finally:
            # Close PDF document
            if pdf_document is not None:
                try:
                    pdf_document.close()
                except:
                    pass
            
            # Clean up the temporary PDF file
            try:
                if os.path.exists(pdf_path):
                    os.unlink(pdf_path)
            except:
                pass
    
    def _generate_docx_visual_preview(self) -> Dict[str, Any]:
        """Create a visual representation of DOCX content using PIL."""
        try:
            doc = DocxDocument(self.file_path)
            
            # Extract text content
            content_lines = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    # Wrap long lines
                    if len(text) > 60:
                        words = text.split()
                        current_line = ""
                        for word in words:
                            if len(current_line + " " + word) <= 60:
                                current_line += (" " if current_line else "") + word
                            else:
                                if current_line:
                                    content_lines.append(current_line)
                                current_line = word
                        if current_line:
                            content_lines.append(current_line)
                    else:
                        content_lines.append(text)
                    
                    # Limit to reasonable number of lines
                    if len(content_lines) >= 15:
                        break
            
            if not content_lines:
                content_lines = ["Document appears to be empty or has no readable text."]
            
            # Create image
            img_width = 400
            line_height = 20
            padding = 20
            img_height = max(250, len(content_lines) * line_height + padding * 2)
            
            # Create a white background image
            img = Image.new('RGB', (img_width, img_height), color='white')
            
            # Try to get a font, fall back to default if not available
            draw = ImageDraw.Draw(img)
            
            # Try to use a system font
            try:
                font = ImageFont.truetype("arial.ttf", 12)
            except:
                try:
                    font = ImageFont.truetype("/System/Library/Fonts/Arial.ttf", 12)  # macOS
                except:
                    try:
                        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 12)  # Linux
                    except:
                        font = ImageFont.load_default()
            
            # Draw text lines
            y_position = padding
            for line in content_lines:
                if y_position + line_height <= img_height - padding:
                    draw.text((padding, y_position), line, fill='black', font=font)
                    y_position += line_height
                else:
                    # Add truncation indicator
                    draw.text((padding, y_position), "...", fill='black', font=font)
                    break
            
            # Convert to base64
            buffer = io.BytesIO()
            img.save(buffer, format='PNG')
            img_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            return {
                "type": "image",
                "data": f"data:image/png;base64,{img_base64}",
                "metadata": {
                    "title": "DOCX Content Preview",
                    "method": "visual_text",
                    "lines_shown": len(content_lines)
                }
            }
            
        except Exception as e:
            raise Exception(f"Visual preview generation failed: {str(e)}")
    
    def _generate_docx_text_preview(self) -> Dict[str, Any]:
        """Generate a simple text preview as fallback."""
        try:
            doc = DocxDocument(self.file_path)
            
            # Extract first few paragraphs
            paragraphs = []
            for i, paragraph in enumerate(doc.paragraphs):
                if i >= 5:  # Limit to first 5 paragraphs
                    break
                text = paragraph.text.strip()
                if text:  # Only add non-empty paragraphs
                    paragraphs.append(text)
            
            if not paragraphs:
                return {
                    "type": "text",
                    "data": "DOCX file appears to be empty or has no readable text."
                }
            
            preview_text = '\n\n'.join(paragraphs)
            
            # Limit total length
            if len(preview_text) > 500:
                preview_text = preview_text[:500] + "..."
            
            return {
                "type": "text",
                "data": preview_text,
                "metadata": {
                    "total_paragraphs": len(doc.paragraphs),
                    "paragraphs_shown": len(paragraphs),
                    "method": "text_only"
                }
            }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating DOCX preview: {str(e)}"
            }
    
    def _generate_text_preview(self) -> Dict[str, Any]:
        """Generate preview for text files - show first few lines."""
        try:
            with open(self.file_path, 'r', encoding='utf-8') as file:
                lines = []
                for i, line in enumerate(file):
                    if i >= 20:  # Limit to first 20 lines
                        break
                    lines.append(line.rstrip())
                
                preview_text = '\n'.join(lines)
                
                # Limit total length
                if len(preview_text) > 500:
                    preview_text = preview_text[:500] + "..."
                
                return {
                    "type": "text",
                    "data": preview_text,
                    "metadata": {
                        "lines_shown": len(lines),
                        "encoding": "utf-8"
                    }
                }
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(self.file_path, 'r', encoding='latin-1') as file:
                    content = file.read(500)
                    return {
                        "type": "text",
                        "data": content + ("..." if len(content) == 500 else ""),
                        "metadata": {
                            "encoding": "latin-1"
                        }
                    }
            except Exception as e:
                return {
                    "type": "error",
                    "data": f"Error reading text file: {str(e)}"
                }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error generating text preview: {str(e)}"
            }
    
    def _generate_default_preview(self) -> Dict[str, Any]:
        """Generate default preview for unsupported file types."""
        try:
            file_size = self.file_path.stat().st_size
            file_size_mb = file_size / (1024 * 1024)
            
            return {
                "type": "info",
                "data": f"Preview not available for {self.file_extension} files.\n\nFile size: {file_size_mb:.2f} MB\nFile type: {self.file_extension.upper()}"
            }
        except Exception as e:
            return {
                "type": "error",
                "data": f"Error getting file info: {str(e)}"
            }
    
    def _resize_image_for_preview(self, img: Image.Image, max_width: int = 400) -> Image.Image:
        """Resize image for preview while maintaining aspect ratio."""
        if img.width <= max_width:
            return img
        
        ratio = max_width / img.width
        new_height = int(img.height * ratio)
        return img.resize((max_width, new_height), Image.Resampling.LANCZOS) 