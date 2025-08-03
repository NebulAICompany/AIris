import pythoncom
from docx2pdf import convert
from pathlib import Path
from PIL import Image
import io
import openpyxl
from backend.pipelines.parsers.tools import describe_image, specify_sentence

def ExcelParser(file_path: str, client=None):
        wb_obj = openpyxl.load_workbook(file_path)
        
        content_parts = []
        for sheet in wb_obj.sheetnames:
            content_parts.append(f"--- Sheet: {sheet} ---\n")
            worksheet = wb_obj[sheet]

            # Eğer bu bir normal worksheet ise
            if isinstance(worksheet, openpyxl.worksheet.worksheet.Worksheet):
                for row in worksheet.iter_rows():
                    for cell in row:
                        value = cell.value if cell.value is not None else ""
                        if not value == "":
                            content_parts.append(f"[{cell.row}, {cell.column}] = {value}\n")


        # Tüm content_parts'ı birleştir ve return et
        extracted_text = "".join(content_parts)
        print(f"Excel text extraction completed. Total length: {len(extracted_text)} characters")
        return extracted_text

def ImageParser(file_path: str, client=None):
    image_path = Path(file_path)
    if not image_path.exists():
        print(f"Görsel dosyası bulunamadı: {file_path}")
        return None

    with open(file_path, "rb") as f:
        image_bytes = f.read()

    # Görseli aç (PIL kullanarak)
    image = Image.open(io.BytesIO(image_bytes))

    # images/ klasörü oluşturulmamışsa oluştur
    images_dir = Path(__file__).parent.parent / "images"
    images_dir.mkdir(exist_ok=True)

    # Image reference için sadece dosya adını kullan (uzantısız)
    image_reference = Path(file_path).stem
    
    # Görseli kaydet
    saved_image_path = images_dir / f"{image_reference}.png"
    image.save(saved_image_path, format='PNG')

    # Açıklamayı al
    description = describe_image(image_bytes)
    updated_description = specify_sentence(description, f"((Image):{image_reference})")

    # Text content'i hazırla
    content = f"Açıklama:\n{updated_description}\n"
    
    print(f"Image text extraction completed. Total length: {len(content)} characters")
    return content


def TxtParser(file_path: str):
        content_parts = []
        with open(file_path, "r", encoding="utf-8") as infile:
            for line in infile:
                if line.strip():  # satır boş değilse
                    content_parts.append(line)

        extracted_text = "".join(content_parts)
        print(f"TXT text extraction completed. Total length: {len(extracted_text)} characters")
        return extracted_text


def DocxParser(file_path: str):
        # Initialize COM at the beginning
        pythoncom.CoInitialize()

        output_dir = "backend/database/uploads"
        output_dir.mkdir(parents=True, exist_ok=True)
        output_pdf = output_dir / (Path(file_path).stem + ".pdf")
        convert(file_path, str(output_pdf))

        file_path = str(output_pdf)
        if not Path(file_path).exists():
            print(f"PDF dosyası bulunamadı: {file_path}")
            return None
        else:
            # PDF'i parse et ve text'i al
            from backend.pipelines.parsers.pdf_parser import PdfParser
            pdf_parser = PdfParser(file_path)
            extracted_text = pdf_parser.run()
            
            print("DOCX to PDF ve text extraction işlemi tamamlandı.")
            
            # Geçici PDF dosyasını sil
            try:
                Path(file_path).unlink()
            except Exception as e:
                print(f"{file_path} silinemedi: {e}")
            
            return extracted_text