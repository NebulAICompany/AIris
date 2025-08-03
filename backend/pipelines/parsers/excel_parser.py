import openpyxl

class ExcelParser:
    def __init__(self, file_path: str):
        self.file_path = file_path
    
    def run(self):
        wb_obj = openpyxl.load_workbook(self.file_path)
        
        content_parts = []
        for sheet in wb_obj.sheetnames:
            content_parts.append(f"--- Sheet: {sheet} ---\n")
            worksheet = wb_obj[sheet]

            # Eğer bu bir normal worksheet ise
            if isinstance(worksheet, openpyxl.worksheet.worksheet.Worksheet):
                for row in worksheet.iter_rows():
                    for cell in row:
                        value = cell.value if cell.value is not None else ""
                        if not value == "":
                            content_parts.append(f"[{cell.row}, {cell.column}] = {value}\n")


        # Tüm content_parts'ı birleştir ve return et
        extracted_text = "".join(content_parts)
        print(f"Excel text extraction completed. Total length: {len(extracted_text)} characters")
        return extracted_text