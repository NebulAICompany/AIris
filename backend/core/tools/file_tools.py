from typing import Any, Dict, List, Optional
from pathlib import Path
import json
import uuid
import base64
import re
from datetime import datetime
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from e2b_code_interpreter import Sandbox
from dotenv import load_dotenv
from backend.shared.logger import get_logger
from backend.shared.constants import CREATED_DOCUMENTS_PATH, REPORTS_CHARTS_FILE, openai_client, IMAGE_EXTENSIONS
from backend.core.prompts import describe_image_prompt
from backend.utils.e2b_file_template import ensure_file_agent_template
from langchain_core.tools import tool 

load_dotenv()

GENERATED_FILES = []

FILES_PATH = Path(CREATED_DOCUMENTS_PATH)
FILES_PATH.mkdir(parents=True, exist_ok=True)

FILE_AGENT_SANDBOX_TEMPLATE = "file-agent-template"

logger = get_logger("FILE_TOOLS")

def _resolve_file_path(file_name: str) -> Path:
    """Resolve a filename to an absolute path within CREATED_DOCUMENTS_PATH.
    Strips any directory components to prevent path traversal."""
    safe_name = Path(file_name).name
    return FILES_PATH / safe_name


def _stage_referenced_files_in_sandbox(code: str, sandbox: Sandbox) -> List[str]:
    """Upload locally created files referenced in code into the sandbox."""
    uploaded_files = []
    pattern = r"""['"]([^'"]+\.(?:pptx|docx|xlsx|xls|png|jpg|jpeg|gif|bmp|pdf|csv|txt))['"]"""
    referenced_paths = re.findall(pattern, code, flags=re.IGNORECASE)

    for referenced_path in referenced_paths:
        local_path = _resolve_file_path(Path(referenced_path).name)
        if not local_path.exists() or not local_path.is_file():
            continue

        sandbox_name = local_path.name
        file_bytes = local_path.read_bytes()
        sandbox.files.write(f"/home/user/{sandbox_name}", file_bytes)
        sandbox.files.write(f"/mnt/data/{sandbox_name}", file_bytes)
        uploaded_files.append(sandbox_name)

    return sorted(set(uploaded_files))

@tool(parse_docstring=True)
def read_excel_file(
    file_name: str, sheet_name: Optional[str] = None, max_rows: int = 100
) -> Dict[str, Any]:
    """Read the contents of an Excel file from the documents directory.

    Args:
        file_name: Name of the Excel file to read (must be in the created documents folder).
        sheet_name: Name of the sheet to read (None reads the active sheet).
        max_rows: Maximum number of data rows to return (default 100).
    """
    try:
        file_path = _resolve_file_path(file_name)
        if not file_path.exists():
            return {"success": False, "error": f"File not found: {file_name}"}

        wb = load_workbook(str(file_path), data_only=True)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return {
                    "success": False,
                    "error": f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(wb.sheetnames)}",
                }
            ws = wb[sheet_name]
        else:
            ws = wb.active

        rows = []
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx >= max_rows:
                break
            rows.append([str(cell) if cell is not None else "" for cell in row])

        total_rows = ws.max_row or 0
        truncated = total_rows > max_rows

        return {
            "success": True,
            "file_name": file_name,
            "sheet_name": ws.title,
            "all_sheets": wb.sheetnames,
            "total_rows": total_rows,
            "rows_returned": len(rows),
            "truncated": truncated,
            "data": rows,
        }

    except Exception as e:
        logger.error(f"Error reading Excel file: {str(e)}")
        return {"success": False, "error": f"Failed to read Excel file: {str(e)}"}


@tool(parse_docstring=True)
def read_word_document(file_name: str) -> Dict[str, Any]:
    """Read the contents of a Word document from the documents directory.

    Args:
        file_name: Name of the Word document to read (must be in the created documents folder).
    """
    try:
        file_path = _resolve_file_path(file_name)
        if not file_path.exists():
            return {"success": False, "error": f"File not found: {file_name}"}

        doc = Document(str(file_path))

        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append({
                    "text": para.text,
                    "style": para.style.name if para.style else "Normal",
                })

        tables = []
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text for cell in row.cells])
            tables.append({
                "table_index": table_idx,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "data": table_rows,
            })

        return {
            "success": True,
            "file_name": file_name,
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "paragraphs": paragraphs,
            "tables": tables,
        }

    except Exception as e:
        logger.error(f"Error reading Word document: {str(e)}")
        return {"success": False, "error": f"Failed to read Word document: {str(e)}"}


@tool(parse_docstring=True)
def read_powerpoint_file(file_name: str) -> Dict[str, Any]:
    """Read the contents of a PowerPoint file from the documents directory.

    Args:
        file_name: Name of the PowerPoint file to read (must be in the created documents folder).
    """
    try:
        file_path = _resolve_file_path(file_name)
        if not file_path.exists():
            return {"success": False, "error": f"File not found: {file_name}"}

        prs = Presentation(str(file_path))

        slides = []
        for slide_idx, slide in enumerate(prs.slides):
            slide_data = {
                "slide_number": slide_idx + 1,
                "layout": slide.slide_layout.name if slide.slide_layout else "Unknown",
                "shapes": [],
                "tables": [],
                "notes": "",
            }

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_content = "\n".join(
                        para.text for para in shape.text_frame.paragraphs if para.text.strip()
                    )
                    if text_content.strip():
                        slide_data["shapes"].append({
                            "name": shape.name,
                            "type": "text",
                            "content": text_content,
                        })

                if shape.has_table:
                    table = shape.table
                    table_rows = []
                    for row in table.rows:
                        table_rows.append([cell.text for cell in row.cells])
                    slide_data["tables"].append({
                        "rows": len(table.rows),
                        "columns": len(table.columns),
                        "data": table_rows,
                    })

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                slide_data["notes"] = slide.notes_slide.notes_text_frame.text

            slides.append(slide_data)

        return {
            "success": True,
            "file_name": file_name,
            "slide_count": len(slides),
            "slides": slides,
        }

    except Exception as e:
        logger.error(f"Error reading PowerPoint file: {str(e)}")
        return {"success": False, "error": f"Failed to read PowerPoint file: {str(e)}"}


@tool(parse_docstring=True)
def describe_file_image(file_name: str, query: str) -> str:
    """Analyze an image file from the documents directory using vision AI.

    Use this tool to understand what is shown in an image (PNG, JPG) that exists
    in the created documents folder.

    Args:
        file_name: Name of the image file to analyze (must be in the created documents folder).
        query: A question or description of what you want to understand about the image.
    """
    try:
        file_path = _resolve_file_path(file_name)
        if not file_path.exists():
            return f"Error: Image file not found: {file_name}"

        ext = file_path.suffix.lower()
        if ext not in IMAGE_EXTENSIONS:
            return f"Error: Unsupported image format '{ext}'. Supported: {', '.join(IMAGE_EXTENSIONS)}"

        with open(file_path, "rb") as f:
            base64_image = base64.b64encode(f.read()).decode("utf-8")

        mime_map = {
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".gif": "image/gif",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
        }
        mime_type = mime_map.get(ext, "image/png")

        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": f"{describe_image_prompt}\n\nUser Query: {query}",
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{mime_type};base64,{base64_image}",
                            "detail": "high",
                        },
                    },
                ],
            }
        ]

        response = openai_client.chat.completions.create(
            model="gpt-5.2",
            messages=messages,
            max_completion_tokens=1500,
            temperature=0.1,
        )

        if not response.choices[0].message.content:
            return "Error: No response received from Vision API"

        analysis = response.choices[0].message.content
        logger.info(f"Successfully analyzed image: {file_name}")
        return f"Image Analysis Results:\n\n{analysis}"

    except Exception as e:
        error_msg = f"Error analyzing image {file_name}: {str(e)}"
        logger.error(error_msg)
        return error_msg


@tool(parse_docstring=True)
def execute_file_code(code: str, output_filenames: List[str]) -> Dict[str, Any]:
    """Execute Python code in a sandboxed environment to create or process files.

    The sandbox has python-pptx, openpyxl, python-docx, and Pillow pre-installed.
    Use this for complex file operations that go beyond what the dedicated tools offer.
    Existing files from the created documents folder that are referenced in the code
    are automatically uploaded into the sandbox at /home/user/<filename> and
    /mnt/data/<filename>. Each output file must be saved under /home/user/ in the sandbox.

    Args:
        code: Complete Python code to execute. Must save output files under /home/user/.
        output_filenames: List of filenames the code will produce under /home/user/ (e.g. ["report.docx", "chart.png"]).
    """
    sandbox = None
    try:
        
        # Check syntax locally before spending time on a sandbox
        try:
            compile(code, "<file_code>", "exec")
        except SyntaxError as e:
            return {
                "success": False,
                "error": (
                    f"SyntaxError before execution: {e.msg} at line {e.lineno}\n"
                    f"  {e.text or ''}"
                    f"Fix the syntax and retry."
                ),
            }

        logger.info(f"Creating sandbox for code execution: {code}")
        try:
            sandbox = Sandbox.create(template=FILE_AGENT_SANDBOX_TEMPLATE, timeout=60)
        except Exception as create_err:
            logger.warning(
                "Sandbox creation failed %s. Building template and retrying once.",
                str(create_err),
            )
            ensure_file_agent_template(alias=FILE_AGENT_SANDBOX_TEMPLATE)
            sandbox = Sandbox.create(template=FILE_AGENT_SANDBOX_TEMPLATE, timeout=60)

        staged_inputs = _stage_referenced_files_in_sandbox(code, sandbox)
        execution = sandbox.run_code(code)

        logger.info(f"Execution result: {execution}")

        if execution.error:
            traceback = execution.error.traceback or ""
            logger.error(f"Code execution error: {execution.error.name}: {execution.error.value}\nTraceback:\n{traceback}\nFix the code and retry.")
            return {
                "success": False,
                "error": (
                    f"Code execution error: {execution.error.name}: {execution.error.value}\n"
                    f"Traceback:\n{traceback}\n"
                    f"Fix the code and retry."
                ),
            }

        saved_files = []
        errors = []

        for fname in output_filenames:
            safe_name = Path(fname).name
            sandbox_path = f"/home/user/{safe_name}"
            suffix = Path(safe_name).suffix
            is_edit_of_existing_file = safe_name in staged_inputs

            if is_edit_of_existing_file:
                local_name = safe_name
                local_path = FILES_PATH / local_name
            else:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                unique_id = uuid.uuid4().hex[:8]
                stem = Path(safe_name).stem
                local_name = f"{stem}_{timestamp}_{unique_id}{suffix}"
                local_path = FILES_PATH / local_name

            try:
                content = sandbox.files.read(sandbox_path, format="bytes")
                with open(local_path, "wb") as f:
                    f.write(content)

                ext = suffix.lower()
                type_map = {
                    ".xlsx": "excel", ".xls": "excel",
                    ".docx": "word", ".doc": "word",
                    ".pptx": "powerpoint",
                    ".png": "image", ".jpg": "image", ".jpeg": "image",
                    ".pdf": "pdf", ".csv": "csv", ".txt": "text",
                }
                file_type = type_map.get(ext, "unknown")

                file_info = {
                    "filename": local_name,
                    "file_path": str(local_path),
                    "file_type": file_type,
                    "created_at": datetime.now().isoformat(),
                    "message": (
                        f"File updated via sandbox: {local_name}"
                        if is_edit_of_existing_file
                        else f"File created via sandbox: {local_name}"
                    ),
                }
                set_generated_files([file_info])
                saved_files.append(file_info)

            except Exception as e:
                errors.append(f"Failed to retrieve '{fname}': {str(e)}")

        stdout = ""
        if execution.logs and execution.logs.stdout:
            stdout = "\n".join(execution.logs.stdout)

        result = {
            "success": len(saved_files) > 0,
            "files_created": saved_files,
            "stdout": stdout,
        }
        if staged_inputs:
            result["staged_input_files"] = staged_inputs
        if errors:
            result["errors"] = errors
        return result

    except Exception as e:
        logger.error(f"Error executing code: {str(e)}")
        return {"success": False, "error": f"Sandbox error: {str(e)}"}
    finally:
        if sandbox:
            try:
                logger.info(f"Killing sandbox: {sandbox}")
                sandbox.kill()
            except Exception as e:
                logger.warning(f"Failed to kill sandbox: {str(e)}")


@tool(parse_docstring=True)
def list_created_files() -> Dict[str, Any]:
    """List all created files in the documents directory.

    Returns a list of all files that have been created, including their names and full paths.
    This helps track which files are available and where they are located.
    """
    try:
        created_files = []

        if FILES_PATH.exists() and FILES_PATH.is_dir():
            for file_path in FILES_PATH.iterdir():
                if file_path.is_file():
                    extension = file_path.suffix.lower()
                    type_map = {
                        ".xlsx": "excel", ".xls": "excel",
                        ".docx": "word", ".doc": "word",
                        ".pptx": "powerpoint",
                        ".pdf": "pdf",
                        ".txt": "text",
                        ".csv": "csv",
                        ".png": "image", ".jpg": "image", ".jpeg": "image",
                        ".gif": "image", ".bmp": "image",
                    }
                    file_type = type_map.get(extension, "unknown")

                    created_files.append({
                        "filename": file_path.name,
                        "file_path": str(file_path),
                        "file_type": file_type,
                        "extension": extension,
                    })

        return {
            "success": True,
            "count": len(created_files),
            "files": created_files,
            "message": f"Found {len(created_files)} created file(s) in the documents directory.",
        }

    except Exception as e:
        logger.error(f"Error listing created files: {str(e)}")
        return {
            "success": False,
            "error": f"Failed to list created files: {str(e)}",
            "files": [],
        }


@tool(parse_docstring=True)
def get_available_charts() -> Dict[str, Any]:
    """Get a list of chart images that have been generated by the plotting agent.

    Returns chart PNG filenames and their descriptions so they can be embedded
    into Word documents or PowerPoint presentations via the sandbox.
    Each chart PNG lives in the created documents folder alongside other files.
    """
    try:
        if not REPORTS_CHARTS_FILE.exists():
            return {
                "success": True,
                "count": 0,
                "charts": [],
                "message": "No charts have been generated yet.",
            }

        with open(REPORTS_CHARTS_FILE, "r", encoding="utf-8") as f:
            records = json.load(f)

        if not isinstance(records, list):
            records = []

        charts = []
        for record in records:
            png_rel = record.get("png_file_path", "")
            png_name = Path(png_rel).name if png_rel else ""
            png_full = FILES_PATH / png_name if png_name else None

            charts.append({
                "filename": png_name,
                "exists": png_full.exists() if png_full else False,
                "description": record.get("description", ""),
                "created_at": record.get("creation_time", ""),
            })

        return {
            "success": True,
            "count": len(charts),
            "charts": charts,
            "message": f"Found {len(charts)} chart(s) available for embedding.",
        }

    except Exception as e:
        logger.error(f"Error reading chart reports: {str(e)}")
        return {
            "success": False,
            "error": f"Failed to read chart reports: {str(e)}",
            "charts": [],
        }


def get_generated_files() -> Dict[str, Any]:
    global GENERATED_FILES
    return GENERATED_FILES


def clear_generated_files():
    global GENERATED_FILES
    GENERATED_FILES.clear()


def set_generated_files(files: List[Dict[str, Any]]):
    global GENERATED_FILES
    for new_file in files:
        GENERATED_FILES = [
            f for f in GENERATED_FILES
            if f.get("filename") != new_file.get("filename")
        ]
        GENERATED_FILES.append(new_file)
